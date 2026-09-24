"""
build_tracking_sheet.py
output/ に蓄積された週次スクリーニング結果・パフォーマンス追跡データを
1つのExcelブック(output/tracking_sheet.xlsx)にまとめる。

Notion連携が使えない代わりに、これを「定点観測の記録簿」として使う。
output/ 配下の全CSVから毎回まるごと作り直す(このスクリプト自身は状態を
持たない)ので、CSVさえ残っていれば何度でも再生成できる。

シート構成:
  - 週次ピック: 毎週の選出銘柄(選出日・順位・スコア・入場株価)を全期間分
  - 価格推移ログ: track_performance.py --check の結果を全期間分、縦持ちで蓄積
  - 価格推移(ピボット): 銘柄 x 確認日 のマトリクスで株価推移を一覧できる形

使い方:
  python3 build_tracking_sheet.py
"""

import logging
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

import config

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
logger = logging.getLogger(__name__)

OUTPUT_DIR = Path(config.OUTPUT_DIR)
PERF_DIR = OUTPUT_DIR / "performance"
XLSX_PATH = OUTPUT_DIR / "tracking_sheet.xlsx"

HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
BODY_FONT = Font(name="Arial")


def _to_date(s):
    return datetime.strptime(str(s), "%Y-%m-%d").date()


def load_weekly_picks() -> pd.DataFrame:
    """全期間の top_candidates_*.csv + entry_*.csv を結合して週次ピック一覧を作る。"""
    rows = []
    for entry_path in sorted(PERF_DIR.glob("entry_*.csv")):
        date_str = entry_path.stem.replace("entry_", "")
        entry_df = pd.read_csv(entry_path, dtype={"code": str})

        cand_path = OUTPUT_DIR / f"top_candidates_{date_str}.csv"
        score_map = {}
        if cand_path.exists():
            cand_df = pd.read_csv(cand_path, dtype={"code": str})
            score_map = cand_df.set_index("code")[["score_valuation", "score_quality"]].to_dict("index")

        for _, r in entry_df.iterrows():
            scores = score_map.get(r["code"], {})
            rows.append({
                "選出日": _to_date(date_str),
                "証券コード": r["code"],
                "銘柄名": r["name"],
                "セクター": r.get("sector", "–"),
                "順位": int(r["rank"]),
                "総合スコア": round(float(r["total_score"]), 1),
                "バリュエーションスコア": round(scores.get("score_valuation", float("nan")), 1) if scores else None,
                "収益性スコア": round(scores.get("score_quality", float("nan")), 1) if scores else None,
                "入場株価": r["entry_price"],
                "TOPIX入場値": r["topix_entry"],
            })
    return pd.DataFrame(rows)


def load_price_log() -> pd.DataFrame:
    """全期間の result_*_checked_*.csv を縦持ちで結合する。"""
    rows = []
    for result_path in sorted(PERF_DIR.glob("result_*_checked_*.csv")):
        df = pd.read_csv(result_path, dtype={"code": str})
        for _, r in df.iterrows():
            rows.append({
                "選出日": _to_date(r["entry_date"]),
                "証券コード": r["code"],
                "銘柄名": r["name"],
                "確認日": _to_date(r["check_date"]),
                "入場株価": r["entry_price"],
                "現在株価": r["current_price"],
                "TOPIXリターン%": (r["topix_return_pct"] / 100.0) if pd.notna(r.get("topix_return_pct")) else None,
            })
    if not rows:
        return pd.DataFrame(columns=["選出日", "証券コード", "銘柄名", "確認日", "入場株価", "現在株価", "TOPIXリターン%"])
    # 同一(選出日,コード,確認日)の重複(再実行等)は最後の行を優先
    df = pd.DataFrame(rows)
    df = df.drop_duplicates(subset=["選出日", "証券コード", "確認日"], keep="last")
    return df.sort_values(["選出日", "証券コード", "確認日"])


def _autosize(ws, df: pd.DataFrame):
    for i, col in enumerate(df.columns, start=1):
        width = max(len(str(col)), *(len(str(v)) for v in df[col].astype(str).tolist() or [""]))
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 40)


def _write_header(ws, headers):
    for i, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=i, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"


def write_weekly_picks_sheet(wb: Workbook, picks: pd.DataFrame):
    ws = wb.create_sheet("週次ピック")
    if picks.empty:
        _write_header(ws, ["選出日", "証券コード", "銘柄名", "セクター", "順位", "総合スコア",
                            "バリュエーションスコア", "収益性スコア", "入場株価", "TOPIX入場値"])
        return
    _write_header(ws, list(picks.columns))
    for r_i, row in enumerate(picks.itertuples(index=False), start=2):
        for c_i, val in enumerate(row, start=1):
            cell = ws.cell(row=r_i, column=c_i, value=val)
            cell.font = BODY_FONT
            col_name = picks.columns[c_i - 1]
            if col_name == "選出日":
                cell.number_format = "yyyy-mm-dd"
            elif col_name in ("入場株価", "TOPIX入場値"):
                cell.number_format = "#,##0.00"
    _autosize(ws, picks)


def write_price_log_sheet(wb: Workbook, log: pd.DataFrame):
    ws = wb.create_sheet("価格推移ログ")
    headers = ["選出日", "証券コード", "銘柄名", "確認日", "経過日数", "入場株価",
               "現在株価", "リターン%", "TOPIXリターン%", "アルファ%pt"]
    _write_header(ws, headers)
    if log.empty:
        return
    for r_i, row in enumerate(log.itertuples(index=False), start=2):
        entry_date, code, name, check_date, entry_price, current_price, topix_return = row
        elapsed_days = (check_date - entry_date).days
        return_pct = (current_price - entry_price) / entry_price if entry_price else None
        alpha_pct = (return_pct - topix_return) if (return_pct is not None and topix_return is not None) else None

        ws.cell(row=r_i, column=1, value=entry_date).number_format = "yyyy-mm-dd"
        ws.cell(row=r_i, column=2, value=code)
        ws.cell(row=r_i, column=3, value=name)
        ws.cell(row=r_i, column=4, value=check_date).number_format = "yyyy-mm-dd"
        ws.cell(row=r_i, column=5, value=elapsed_days)
        ws.cell(row=r_i, column=6, value=entry_price).number_format = "#,##0.00"
        ws.cell(row=r_i, column=7, value=current_price).number_format = "#,##0.00"
        ws.cell(row=r_i, column=8, value=return_pct).number_format = "0.0%"
        ws.cell(row=r_i, column=9, value=topix_return).number_format = "0.0%"
        ws.cell(row=r_i, column=10, value=alpha_pct).number_format = "0.0%"
        for c_i in range(1, 11):
            ws.cell(row=r_i, column=c_i).font = BODY_FONT
    for i, w in enumerate([12, 10, 20, 12, 10, 12, 12, 10, 12, 10], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_pivot_sheet(wb: Workbook, picks: pd.DataFrame, log: pd.DataFrame):
    ws = wb.create_sheet("価格推移(ピボット)")
    base_headers = ["選出日", "証券コード", "銘柄名", "入場株価"]
    check_dates = sorted(log["確認日"].unique()) if not log.empty else []

    _write_header(ws, base_headers + [d.strftime("%Y-%m-%d") for d in check_dates])
    for i, d in enumerate(check_dates, start=len(base_headers) + 1):
        ws.cell(row=1, column=i).number_format = "yyyy-mm-dd"

    if picks.empty:
        return

    stocks = picks[["選出日", "証券コード", "銘柄名", "入場株価"]].drop_duplicates(
        subset=["選出日", "証券コード"]
    ).sort_values(["選出日", "証券コード"])

    # (選出日, 証券コード, 確認日) -> 現在株価 のルックアップ表(Pythonで事前計算し、値として書き込む)
    price_lookup = {}
    if not log.empty:
        for row in log.itertuples(index=False):
            entry_date, code, _name, check_date, _entry_price, current_price, _topix = row
            price_lookup[(entry_date, code, check_date)] = current_price

    for r_i, row in enumerate(stocks.itertuples(index=False), start=2):
        entry_date, code, name, entry_price = row
        ws.cell(row=r_i, column=1, value=entry_date).number_format = "yyyy-mm-dd"
        ws.cell(row=r_i, column=2, value=code)
        ws.cell(row=r_i, column=3, value=name)
        ws.cell(row=r_i, column=4, value=entry_price).number_format = "#,##0.00"
        for c_i in range(1, 5):
            ws.cell(row=r_i, column=c_i).font = BODY_FONT
        for d_i, d in enumerate(check_dates, start=len(base_headers) + 1):
            price = price_lookup.get((entry_date, code, d))
            cell = ws.cell(row=r_i, column=d_i, value=price)
            cell.number_format = "#,##0.00"
            cell.font = BODY_FONT

    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 12
    for i in range(len(base_headers) + 1, len(base_headers) + 1 + len(check_dates)):
        ws.column_dimensions[get_column_letter(i)].width = 12


def build() -> Path:
    picks = load_weekly_picks()
    log = load_price_log()

    wb = Workbook()
    wb.remove(wb.active)  # デフォルトの空シートを削除
    write_weekly_picks_sheet(wb, picks)
    write_price_log_sheet(wb, log)
    write_pivot_sheet(wb, picks, log)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(XLSX_PATH)
    logger.info(f"トラッキングシートを出力しました: {XLSX_PATH}")
    return XLSX_PATH


if __name__ == "__main__":
    build()
